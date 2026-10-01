# -*- coding: utf-8 -*-
"""build_monitoreo_corredor.py — Traducción de carga difusa a concentración
por subtramo (Actividad 5, monitoreo del corredor biológico).

    cd scripts/tramos && node build_posiciones_monitoreo.mjs   # una vez
    python scripts/build_monitoreo_corredor.py

Procedimiento (definido por la dirección del proyecto):
  1. Universo: subtramos en categoría de carga de N Alta o Muy alta (Tabla
     5.2, caña actualizada), excepto los ríos descartados de entrada: Nima,
     Guachal y Zabaletas (sin caudal en las fuentes de estación) y Risaralda
     (sin N ni P total).
  2. Estación de calidad: la de cierre del subtramo, que es la más aguas
     abajo del tramo, siempre que tenga registros de NT y PT.
  3. Carga = ha de caña × dosis × fracción de pérdida × factor de retención
     del corredor (1 / 0,8 / 0,6), en kg/d. El factor de emisión del Cap. 4
     ya es dosis × fracción (9,568 kg N y 1,376 kg P /ha/año), así que la
     fracción NO se aplica dos veces.
  4. Concentración = carga / (Q × 86,4), con Q = caudal de referencia de la
     estación hidrométrica más cercana sobre el eje del mismo río.
  5. Contraste con la concentración histórica media (mismo criterio que
     data/calidad_agua/estadisticas_puntos.csv: "<X" cuenta como X).
  6. Dominancia de carga puntual: Tabla 4.10 para Bolo y Fraile. Para el
     resto se suma la carga de DBO5 y SST del inventario de tasas
     retributivas de la CVC (docs/Tasas retributivas.xlsx) por subtramo, sin
     convertirla a N ni a P: no hay un factor validado en el proyecto, así
     que esos subtramos quedan Pendiente con la comparación a la vista.

Asignación de vertimientos: un vertimiento pertenece al subtramo si su
"estación aguas abajo" (llevada a la red de calidad con
data/fuentes/monitoreo_corredor/equivalencia_estaciones_tasas.csv) queda
entre la estación de inicio (exclusiva) y la de cierre (inclusiva) sobre el
eje. Las coordenadas (EPSG:9377) solo se usan como control.

Produce:
    docs/monitoreo_corredor.xlsx               (fórmulas vivas, para verificar)
    data/cartografia/Monitoreo_corredor.geojson (capa "Monitoreo de calidad de agua")
    src/config/monitoreo.js                    (constantes del geovisor, generado:
                                                 mismas cifras que la hoja Parámetros)

Los valores del GeoJSON y de la hoja Verificación se calculan aquí con los
mismos parámetros por defecto que el libro; si difieren de las fórmulas, algo
se desalineó entre ambos.
"""

import csv
import datetime
import json
import math
import os
import sys
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

try:
    sys.stdout.reconfigure(encoding='utf-8')
except (AttributeError, ValueError):
    pass

SCRIPT_DIR  = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(SCRIPT_DIR)
sys.path.insert(0, SCRIPT_DIR)
from build_calidad_trib import to_num  # noqa: E402  — mismo criterio que estadisticas_puntos.csv


def ruta(*partes):
    return os.path.join(PROJECT_DIR, *partes)


PRIORIZACION_CSV = ruta('data', 'fuentes', 'priorizacion', 'Priorizacion_NP_subtramos.csv')
TRAMOS_CSV       = ruta('docs', 'tramos_cana_tributarios.csv')
POSICIONES_CSV   = ruta('data', 'fuentes', 'monitoreo_corredor', 'posiciones_estaciones.csv')
TABLA_4_10_CSV   = ruta('data', 'fuentes', 'monitoreo_corredor', 'tabla_4_10_carga_puntual.csv')
PUNTOS_GEOJSON   = ruta('data', 'calidad_agua', 'puntos_calidad_tributarios.geojson')
CSV_POR_PUNTO    = ruta('data', 'calidad_agua', 'csv_por_punto')
HIDRO_JSON       = ruta('data', 'hidrologia', 'estaciones_hidro_trib.json')
TASAS_XLSX       = ruta('docs', 'Tasas retributivas.xlsx')
EQUIV_TASAS_CSV  = ruta('data', 'fuentes', 'monitoreo_corredor', 'equivalencia_estaciones_tasas.csv')
POLIGONOS_TRAMOS = ruta('data', 'cartografia', 'Priorizacion_NP_tramos.geojson')
EJES_TRIBUTARIOS = ruta('data', 'cartografia', 'Tributarios_rios_cauca.geojson')
SERIES_DIR       = ruta('data', 'hidrologia', 'tributarios')

XLSX_OUT    = ruta('docs', 'monitoreo_corredor.xlsx')
GEOJSON_OUT = ruta('data', 'cartografia', 'Monitoreo_corredor.geojson')
CONFIG_JS_OUT = ruta('src', 'config', 'monitoreo.js')

# ── Parámetros (Informe Fase 1, Cap. 4) ────────────────────────────────
DOSIS_N   = 184.0      # kg N/ha/año — Tabla 4.6 (Narváez, 2024)
DOSIS_P2O5 = 46.0      # kg P2O5/ha/año — Tabla 4.6
CONV_P2O5_P = 0.436    # 2 × 30,97 / 141,94 — Tabla 4.6
FRAC_N    = 0.052      # 5,20 % — Tabla 4.7, escenario medio (Scarpare et al., 2023)
FRAC_P    = 0.0686     # 6,86 % — Tabla 4.8, escenario medio (Jeong et al., 2011)
RETENCION = {'base': 1.0, 'esc1': 0.8, 'esc2': 0.6}
DIAS_ANO  = 365
CONV_Q    = 86.4       # kg/d ÷ (m³/s × 86,4) = mg/L
UMBRAL_DOMINANCIA = 0.5
CONDICIONES = ['Promedio', 'Verano', 'Transición', 'Invierno']
CONDICION_DEFECTO = 'Promedio'

CATEGORIAS_UNIVERSO = ('Muy alta', 'Alta')
MIN_DIAS_SERIE = 5 * 365   # serie hidrométrica mínima para asignar caudal
DEDUPE_KM = 0.15           # mismo punto físico registrado dos veces (segmentacion.mjs)

EXCLUSION_SIN_CAUDAL = ('Sin datos de caudal en las fuentes de estación '
                        '(CVC histórico / portal hidroclimatológico)')
EXCLUIDOS_ENTRADA = {
    'nima':      EXCLUSION_SIN_CAUDAL,
    'guachal':   EXCLUSION_SIN_CAUDAL,
    'zabaletas': EXCLUSION_SIN_CAUDAL,
    'risaralda': 'Sin monitoreo de nitrógeno total y fósforo total',
}

CRITERIO_CON_VERTIMIENTOS = ('Pendiente: DBO5/SST sin factor de conversión a N validado; '
                             'requiere definición de la dirección del proyecto')
CRITERIO_SIN_VERTIMIENTOS = 'Sin vertimientos puntuales digitalizados en el tramo'
CRITERIO_N_MANUAL = 'Aproximado: carga puntual de N equivalente ingresada en el bloque 3'
# Cuencas del inventario que corresponden a los subtramos pendientes
CUENCAS_TASAS = ('Amaime', 'Bugalagrande', 'Desbaratado', 'Guabas', 'La Paila', 'Tulua')

# ── Estilo del proyecto: dos colores, Arial, bordes finos negros ───────
AZUL_PETROLEO = '1F4E5F'
AZUL_CLARO    = 'DCE6EA'
F_BASE   = Font(name='Arial', size=10)
F_BOLD   = Font(name='Arial', size=10, bold=True)
F_TITULO = Font(name='Arial', size=12, bold=True)
F_HEADER = Font(name='Arial', size=10, bold=True, color='FFFFFF')
F_NOTA   = Font(name='Arial', size=9, italic=True)
FILL_HEADER = PatternFill('solid', fgColor=AZUL_PETROLEO)
FILL_INPUT  = PatternFill('solid', fgColor=AZUL_CLARO)
_fino = Side(style='thin', color='000000')
BORDE = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
WRAP_TOP = Alignment(wrap_text=True, vertical='top')
CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)

FMT_HA   = '#,##0.00'
FMT_KGD  = '#,##0.00'
FMT_Q    = '0.000'
FMT_CONC = '0.0000'
FMT_PCT  = '0.0%'
FMT_RAT  = '0.000'


# ── Utilidades ─────────────────────────────────────────────────────────
def normalizar_rio(nombre):
    """Misma clave que normalizeRiver() de src/tramos/stations.js."""
    s = unicodedata.normalize('NFD', str(nombre or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn').lower().strip()
    if s.startswith('rio '):
        s = s[4:]
    s = ' '.join(s.split())
    return {'frayle': 'fraile', 'sabaletas': 'zabaletas'}.get(s, s)


def leer_csv(path, encoding='utf-8-sig'):
    with open(path, encoding=encoding, newline='') as fh:
        return [r for r in csv.DictReader(fh)
                if any(isinstance(v, str) and v and not v.startswith('#') for v in r.values())]


RIO_DISPLAY = {'Tulua': 'Tuluá', 'Riofrio': 'Riofrío'}


def rio_display(rio):
    """'Rio Tulua' → 'Río Tuluá', para títulos y etiquetas del mapa."""
    nombre = rio[4:] if rio.startswith('Rio ') else rio
    return 'Río ' + RIO_DISPLAY.get(nombre, nombre)


def fmt_es(x, d=2):
    """Número con coma decimal y punto de miles (registro de los informes)."""
    return f'{x:,.{d}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def fmt_km(x):
    return f'{x:,.1f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def media(vals):
    return sum(vals) / len(vals) if vals else None


# ── Carga de insumos ───────────────────────────────────────────────────
def cargar_insumos():
    prior = leer_csv(PRIORIZACION_CSV)
    tramos = {(normalizar_rio(r['rio']), int(r['tramo'])): r
              for r in leer_csv(TRAMOS_CSV) if r.get('tramo', '').isdigit()}
    posiciones = leer_csv(POSICIONES_CSV)
    tabla_4_10 = leer_csv(TABLA_4_10_CSV)
    with open(PUNTOS_GEOJSON, encoding='utf-8') as fh:
        puntos = {f['properties']['Punto_Monitoreo']: f for f in json.load(fh)['features']}
    with open(HIDRO_JSON, encoding='utf-8') as fh:
        hidro = {h['nombre']: h for h in json.load(fh)}
    return prior, tramos, posiciones, tabla_4_10, puntos, hidro


def historico(punto_feature):
    """Media, n de NT y PT de una estación desde su CSV por punto."""
    path = os.path.join(CSV_POR_PUNTO, punto_feature['properties']['csv_filename'])
    filas = leer_csv(path)
    out = {}
    for clave, col in (('nt', 'Nitrógeno Total (mg N/l)'), ('pt', 'Fósforo Total (mg P/l)')):
        vals = [v for v in (to_num(r.get(col)) for r in filas) if v is not None]
        out[clave] = round(media(vals), 4) if vals else None
        out[f'n_{clave}'] = len(vals)
    return out


def serie_caudal(nombre):
    path = os.path.join(SERIES_DIR, nombre, 'caudal_diario.csv')
    serie = []
    for r in leer_csv(path):
        try:
            q = float(r['CAUDAL_M3S'])
        except (TypeError, ValueError):
            continue
        if math.isnan(q):
            continue
        serie.append((datetime.date.fromisoformat(r['FECHA']), q))
    return serie


def caudales_por_condicion(serie, umbral_verano, umbral_invierno):
    qs = [q for _, q in serie]
    grupos = {
        'Promedio':   qs,
        'Verano':     [q for q in qs if q <= umbral_verano],
        'Transición': [q for q in qs if umbral_verano < q < umbral_invierno],
        'Invierno':   [q for q in qs if q >= umbral_invierno],
    }
    return {c: media(v) for c, v in grupos.items()}, {c: len(v) for c, v in grupos.items()}


# ── Selección de estaciones ────────────────────────────────────────────
def seleccionar(prior, tramos, posiciones, puntos, hidro):
    pos_cal = {}
    pos_hid = {}
    for p in posiciones:
        destino = pos_cal if p['tipo'] == 'calidad' else pos_hid
        destino.setdefault(p['clave_rio'], []).append(p)

    universo = [r for r in prior if r['categoria'] in CATEGORIAS_UNIVERSO]
    ejercicio, previas = [], []

    for r in universo:
        clave = normalizar_rio(r['rio'])
        tramo = tramos[(clave, int(r['tramo']))]
        base = {
            'clave': clave,
            'rio': r['rio'],
            'subtramo': int(r['tramo']),
            'posicion_tabla_5_2': r['n_tabla'],
            'categoria': r['categoria'],
            'delimitacion': f"{tramo['estacion_aguas_arriba']} → {tramo['estacion_aguas_abajo']}",
            'estacion_inicio': tramo['estacion_aguas_arriba'],
            'estacion_cierre': tramo['estacion_aguas_abajo'],
            'area_cana_ha': float(tramo['cana_ha_normalizada']),
        }
        if clave in EXCLUIDOS_ENTRADA:
            base['estacion'] = tramo['estacion_aguas_abajo']
            base['motivo_exclusion'] = 'Exclusión previa del río: ' + EXCLUIDOS_ENTRADA[clave]
            previas.append(base)
            continue

        # Estación de calidad: la de cierre si tiene NT y PT; si no, la
        # siguiente aguas abajo que los tenga ("o inmediatamente después").
        cands = sorted(pos_cal[clave], key=lambda p: float(p['km_eje']))
        cierre = next(p for p in cands if p['nombre'] == tramo['estacion_aguas_abajo'])
        km_cierre = float(cierre['km_eje'])
        elegida, hist = None, None
        # La de cierre va primero: un duplicado del mismo punto físico (mismo
        # km) solo entra si la de cierre no tiene NT o PT.
        siguientes = [c for c in cands if c is not cierre and float(c['km_eje']) >= km_cierre - DEDUPE_KM]
        for p in [cierre] + siguientes:
            h = historico(puntos[p['nombre']])
            if h['n_nt'] > 0 and h['n_pt'] > 0:
                elegida, hist = p, h
                break
        if elegida is None:
            raise SystemExit(f"{r['rio']} {r['tramo']}: ninguna estación con NT y PT aguas abajo")
        km_q = float(elegida['km_eje'])

        motivo = ('Estación de cierre del subtramo (la más aguas abajo del tramo)'
                  if elegida is cierre else
                  'Primera estación aguas abajo del cierre con registros de NT y PT')
        motivo += f"; registros NT n={hist['n_nt']}, PT n={hist['n_pt']}"
        duplicados = [c['nombre'] for c in cands if c is not elegida
                      and abs(float(c['km_eje']) - km_q) < DEDUPE_KM]
        for d in duplicados:
            motivo += (f'. Se descarta «{d}»: mismo punto físico (km {fmt_km(km_q)}), '
                       'registro con menos campañas')

        # Estación hidrométrica: la más cercana sobre el eje del mismo río,
        # con al menos cinco años de serie.
        hs = [h for h in pos_hid.get(clave, []) if int(h['n_dias'] or 0) >= MIN_DIAS_SERIE]
        descartadas = [h for h in pos_hid.get(clave, []) if h not in hs]
        h = min(hs, key=lambda h: abs(float(h['km_eje']) - km_q))
        d = float(h['km_eje']) - km_q
        if abs(d) < 0.5:
            relacion = f'Coincidente ({fmt_km(abs(d))} km)'
        elif d > 0:
            relacion = f'{fmt_km(d)} km aguas abajo'
        else:
            relacion = f'{fmt_km(-d)} km aguas arriba'
        meta = hidro[h['nombre']]
        notas = []
        if h['en_zona_canera'] == 'no':
            notas.append('fuera de la zona cañera')
        if meta['estado'] == 'Suspendida':
            notas.append(f"suspendida; serie {meta['años_datos']}")
        elif not meta['años_datos'].endswith(('2025', '2026')):
            notas.append(f"serie {meta['años_datos']}")
        for x in descartadas:
            notas.append(f"se descarta {x['nombre']} ({x['n_dias']} días, {x['anios_datos']})")
        if notas:
            relacion += ' — ' + '; '.join(notas)

        base.update({
            'estacion': elegida['nombre'],
            'km_estacion': round(km_q, 2),
            'motivo_seleccion': motivo,
            'lon': float(elegida['lon']), 'lat': float(elegida['lat']),
            'estacion_hidro': h['nombre'],
            'km_hidro': round(float(h['km_eje']), 2),
            'relacion_hidro': relacion,
            **hist,
        })
        ejercicio.append(base)

    # Correspondencia 1:1 estación ↔ subtramo
    usadas = {}
    for e in ejercicio:
        usadas.setdefault(e['estacion'], []).append(f"{e['rio']} {e['subtramo']}")
    compartidas = {k: v for k, v in usadas.items() if len(v) > 1}

    ejercicio.sort(key=lambda e: (e['clave'], e['subtramo']))
    previas.sort(key=lambda e: (e['clave'], e['subtramo']))
    return ejercicio, previas, compartidas


# ── Inventario de tasas retributivas (CVC) ─────────────────────────────
def limpiar(texto):
    """Nombre comparable: sin espacios duros ni dobles, sin bordes."""
    return ' '.join(str(texto or '').replace('\xa0', ' ').split())


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def control_espacial():
    """Localizador de subtramo por coordenadas EPSG:9377. Devuelve None si
    faltan pyproj/shapely: el control es informativo, no decide nada."""
    try:
        import pyproj
        from shapely.geometry import Point, shape
    except ImportError:
        return None
    tr = pyproj.Transformer.from_crs('EPSG:9377', 'EPSG:4326', always_xy=True)
    with open(POLIGONOS_TRAMOS, encoding='utf-8') as fh:
        poligonos = [(f"{f['properties']['rio']} {f['properties']['tramo']}", shape(f['geometry']))
                     for f in json.load(fh)['features']]

    def localizar(norte, este):
        if norte is None or este is None:
            return None, None, 'sin coordenadas'
        lon, lat = tr.transform(este, norte)
        dentro = [n for n, g in poligonos if g.contains(Point(lon, lat))]
        return lon, lat, ' / '.join(dentro)
    return localizar


def cargar_vertimientos(posiciones):
    """Vertimientos del inventario, con su estación equivalente en la red de
    calidad y su km sobre el eje (si la estación es del cauce principal)."""
    from openpyxl import load_workbook
    equiv = {limpiar(r['estacion_inventario']): r for r in leer_csv(EQUIV_TASAS_CSV)}
    km = {(p['clave_rio'], p['nombre']): float(p['km_eje']) for p in posiciones if p['tipo'] == 'calidad'}
    localizar = control_espacial()

    ws = load_workbook(TASAS_XLSX, data_only=True).worksheets[0]
    encab = [limpiar(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    esperado = ['Vertimiento', 'Tipo de vertimiento', 'Cuerpo de agua receptor del vertimiento',
                'Tramo del cuerpo de agua receptor del vertimiento',
                'Nombre de la estación aguas abajo del vertimiento', 'Cuenca']
    if encab[:6] != esperado:
        raise SystemExit(f'Encabezado inesperado en {TASAS_XLSX}:\n  {encab}')

    out = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not any(fila):
            continue
        v = {
            'registro': limpiar(fila[0]), 'tipo': limpiar(fila[1]), 'receptor': limpiar(fila[2]),
            'tramo_inv': limpiar(fila[3]), 'estacion_inv': limpiar(fila[4]), 'cuenca': limpiar(fila[5]),
            'norte': num(fila[6]), 'este': num(fila[7]), 'municipio': limpiar(fila[8]),
            'q_l_s': num(fila[9]) or 0.0, 'dbo_kg_ano': num(fila[12]) or 0.0, 'sst_kg_ano': num(fila[13]) or 0.0,
        }
        if localizar:
            v['lon'], v['lat'], v['control'] = localizar(v['norte'], v['este'])
        else:
            v['lon'], v['lat'], v['control'] = None, None, ''
        eq = equiv.get(v['estacion_inv'])
        if eq:
            v['estacion_eq'] = eq['estacion_calidad']
            v['clave_rio'] = normalizar_rio(eq['rio'])
            v['km'] = km.get((v['clave_rio'], eq['estacion_calidad']))
            if v['km'] is None:
                raise SystemExit(f"Equivalencia sin posición sobre el eje: {eq['estacion_calidad']}")
        else:
            v['estacion_eq'], v['clave_rio'], v['km'] = '', '', None
        out.append(v)
    return out


def entregas_por_afluente(pendientes, vertimientos, posiciones):
    """{etiqueta: [vertimientos]} de los no asignados cuyo texto de estación
    indica que desembocan en el río del subtramo y cuya proyección sobre el
    eje cae entre la estación de inicio y la de cierre. Informativo."""
    try:
        from shapely.geometry import Point, shape
    except ImportError:
        return {}
    with open(EJES_TRIBUTARIOS, encoding='utf-8') as fh:
        ejes = {normalizar_rio(f['properties']['NOM1_DRENA']): shape(f['geometry'])
                for f in json.load(fh)['features']}
    pos = {(p['clave_rio'], p['nombre']): Point(float(p['lon']), float(p['lat']))
           for p in posiciones if p['tipo'] == 'calidad'}
    out = {}
    for e in pendientes:
        eje = ejes.get(e['clave'])
        if eje is None:
            continue
        t0 = eje.project(pos[(e['clave'], e['estacion_inicio'])])
        t1 = eje.project(pos[(e['clave'], e['estacion_cierre'])])
        lo, hi = min(t0, t1), max(t0, t1)
        nombre = normalizar_rio(e['rio'])
        for v in vertimientos:
            texto = normalizar_rio(v['estacion_inv'])
            if (v['subtramo'] or v['lon'] is None or 'desembocadura' not in texto
                    or f'rio {nombre}' not in texto or e['etiqueta'] in v['control'].split(' / ')):
                continue
            # Si el receptor es un río con eje propio (p. ej. el Nima), la
            # entrega ocurre en su desembocadura, no donde está el vertimiento.
            punto = Point(v['lon'], v['lat'])
            eje_receptor = ejes.get(normalizar_rio(v['receptor']))
            if eje_receptor is not None and eje_receptor is not eje:
                partes = getattr(eje_receptor, 'geoms', [eje_receptor])
                extremos = [Point(c) for g in partes for c in (g.coords[0], g.coords[-1])]
                punto = min(extremos, key=eje.distance)
            if lo < eje.project(punto) <= hi:
                out.setdefault(e['etiqueta'], []).append(v)
    return out


def asignar_vertimientos(ejercicio, vertimientos, posiciones):
    """Asigna cada vertimiento al subtramo pendiente cuya delimitación
    (inicio exclusivo, cierre inclusivo) contiene su estación aguas abajo.
    Devuelve los vertimientos que se listan en la hoja Vertimientos."""
    km = {(p['clave_rio'], p['nombre']): float(p['km_eje']) for p in posiciones if p['tipo'] == 'calidad'}
    pendientes = [e for e in ejercicio if not e.get('tramo_porh')]
    for e in pendientes:
        e['etiqueta'] = f"{e['rio']} {e['subtramo']}"
    etiquetas = {e['etiqueta'] for e in pendientes}
    for v in vertimientos:
        v['subtramo'] = ''
    for e in pendientes:
        ini = km[(e['clave'], e['estacion_inicio'])]
        fin = km[(e['clave'], e['estacion_cierre'])] + DEDUPE_KM
        for v in vertimientos:
            if v['clave_rio'] == e['clave'] and v['km'] is not None and ini < v['km'] <= fin:
                v['subtramo'] = e['etiqueta']

    # Observaciones del control espacial (no cambian la asignación)
    for v in vertimientos:
        obs = []
        sin_coord = v['control'] == 'sin coordenadas'
        en = [] if sin_coord else [t for t in v['control'].split(' / ') if t]
        if sin_coord:
            obs.append('Sin coordenadas numéricas')
        if v['subtramo'] and en and v['subtramo'] not in en:
            obs.append(f"Coordenadas dentro de {' / '.join(en)}; se mantiene la asignación por estación")
        elif v['subtramo'] and not en and not sin_coord:
            obs.append('Coordenadas fuera de la franja de 700 m')
        dentro_pend = [t for t in en if t in etiquetas and t != v['subtramo']]
        if not v['subtramo'] and dentro_pend:
            obs.append(f"Coordenadas dentro de {' / '.join(dentro_pend)}, sin estación equivalente del cauce "
                       'principal; no se suma')
        elif not v['estacion_eq'] and 'desembocadura' in v['estacion_inv'].lower():
            obs.append('Entrega a un afluente o canal, sin estación del cauce principal; no se suma')
        v['observacion'] = '. '.join(obs)

    # Posibles duplicados: mismas coordenadas, otra estación y caudal casi igual
    for i, a in enumerate(vertimientos):
        for b in vertimientos[i + 1:]:
            if (a['lon'] is not None and b['lon'] is not None and a['estacion_inv'] != b['estacion_inv']
                    and abs(a['lon'] - b['lon']) < 5e-4 and abs(a['lat'] - b['lat']) < 5e-4
                    and a['q_l_s'] > 0 and abs(a['q_l_s'] - b['q_l_s']) / a['q_l_s'] < 0.05):
                for x, y in ((a, b), (b, a)):
                    nota = (f"Mismas coordenadas y caudal casi igual al registro {y['registro']} "
                            f"(estación «{y['estacion_inv']}»): posible duplicado")
                    x['observacion'] = '. '.join(t for t in (x['observacion'], nota) if t)

    # Entregas por afluente o canal: el inventario dice que el receptor
    # desemboca en el río del subtramo, pero no hay estación del cauce
    # principal. Se proyectan sus coordenadas sobre el eje solo para avisar.
    entregas = entregas_por_afluente(pendientes, vertimientos, posiciones)

    for e in pendientes:
        asign = [v for v in vertimientos if v['subtramo'] == e['etiqueta']]
        e['n_vertimientos'] = len(asign)
        e['q_vertido_l_s'] = sum(v['q_l_s'] for v in asign)
        e['carga_dbo5_kg_d'] = sum(v['dbo_kg_ano'] for v in asign) / DIAS_ANO
        e['carga_sst_kg_d'] = sum(v['sst_kg_ano'] for v in asign) / DIAS_ANO
        e['registros_vertimientos'] = [v['registro'] for v in asign]
        e['situacion_puntual'] = 'con_vertimientos' if asign else 'sin_vertimientos_digitalizados'
        e['criterio_dominancia'] = CRITERIO_CON_VERTIMIENTOS if asign else CRITERIO_SIN_VERTIMIENTOS
        e['motivo_exclusion'] = e['criterio_dominancia']
        e['notas_vertimientos'] = [
            f"Registro {v['registro']} ({v['receptor']})"
            + (f", asignado a {v['subtramo']} por estación" if v['subtramo'] and v['subtramo'] != e['etiqueta']
               else '')
            + f": {v['observacion']}" for v in vertimientos
            if v['observacion'] and (v['subtramo'] == e['etiqueta']
                                     or e['etiqueta'] in v['control'].split(' / '))]
        e['notas_vertimientos'] += [
            f"Registro {v['registro']} ({v['receptor']}): entrega al río por un afluente o canal "
            f"(«{v['estacion_inv']}»); su punto de entrega, proyectado sobre el eje, cae en este subtramo. "
            f"No se suma: sin estación del cauce principal ({fmt_es(v['dbo_kg_ano'] / DIAS_ANO)} kg DBO5/d)"
            for v in entregas.get(e['etiqueta'], [])]
    return [v for v in vertimientos
            if v['cuenca'] in CUENCAS_TASAS or v['subtramo']
            or any(t in etiquetas for t in v['control'].split(' / '))]


# ── Cálculo (réplica de las fórmulas del libro) ────────────────────────
def calcular(ejercicio, tabla_4_10, caudales):
    fe_n = DOSIS_N * FRAC_N
    fe_p = DOSIS_P2O5 * CONV_P2O5_P * FRAC_P
    porh = {}
    for r in tabla_4_10:
        porh.setdefault((normalizar_rio(r['rio_geovisor']), int(r['subtramo_geovisor'])), []).append(r)

    for e in ejercicio:
        q = caudales[e['estacion_hidro']]['q'][CONDICION_DEFECTO]
        e['q_ref'] = q
        e['q_cond'] = caudales[e['estacion_hidro']]['q']
        for esc, ret in RETENCION.items():
            cn = e['area_cana_ha'] * fe_n * ret / DIAS_ANO
            cp = e['area_cana_ha'] * fe_p * ret / DIAS_ANO
            e[f'carga_n_{esc}'] = cn
            e[f'carga_p_{esc}'] = cp
            e[f'conc_n_{esc}'] = cn / (q * CONV_Q)
            e[f'conc_p_{esc}'] = cp / (q * CONV_Q)
        e['cociente_n'] = e['conc_n_base'] / e['nt'] if e['nt'] else None
        e['cociente_p'] = e['conc_p_base'] / e['pt'] if e['pt'] else None

        filas = porh.get((e['clave'], e['subtramo']))
        if filas:
            dom = dominancia(filas, e['carga_n_base'], e['carga_p_base'])
            e.update(dom)
            e['estado'] = 'Excluida' if dom['frac_puntual_n'] > UMBRAL_DOMINANCIA else 'Incluida'
            e['motivo_exclusion'] = motivo_exclusion(e) if e['estado'] == 'Excluida' else ''
        else:
            # Criterio y motivo los fija asignar_vertimientos() (DBO5/SST)
            e.update({'frac_puntual_n': None, 'frac_puntual_p': None,
                      'tramo_porh': '', 'correspondencia_porh': ''})
            e['estado'] = 'Pendiente'
    return fe_n, fe_p


def dominancia(filas, fe_carga_n, fe_carga_p):
    """Fracción puntual por nutriente según la regla confirmada:
    máximo entre campañas con difusa QUAL2Kw definida (>0); si ninguna la
    tiene, se contrasta contra la carga difusa por factor de emisión."""
    out = {'tramo_porh': filas[0]['tramo_porh'], 'correspondencia_porh': filas[0]['correspondencia']}
    for nut, clave, fe in (('Nitrógeno total', 'n', fe_carga_n), ('Fósforo total', 'p', fe_carga_p)):
        rs = [r for r in filas if r['nutriente'] == nut]
        definidas = [float(r['carga_puntual_kg_d']) /
                     (float(r['carga_puntual_kg_d']) + float(r['carga_difusa_qual2kw_kg_d']))
                     for r in rs if float(r['carga_difusa_qual2kw_kg_d']) > 0]
        if definidas:
            out[f'frac_puntual_{clave}'] = max(definidas)
            aprox = False
        else:
            out[f'frac_puntual_{clave}'] = max(float(r['carga_puntual_kg_d']) /
                                              (float(r['carga_puntual_kg_d']) + fe) for r in rs)
            aprox = True
        if clave == 'n':
            out['criterio_dominancia'] = (
                'Aproximado: difusa QUAL2Kw no definida en el tramo; contraste con la difusa por factor de emisión'
                if aprox else 'Riguroso: Tabla 4.10 (puntual frente a difusa QUAL2Kw)')
    return out


def motivo_exclusion(e):
    pct = f"{e['frac_puntual_n'] * 100:.1f}".replace('.', ',')
    return f"Carga puntual dominante ({pct} % de la carga total de N; {e['criterio_dominancia']})"


# ── Excel ──────────────────────────────────────────────────────────────
def celda(ws, ref, valor, fmt=None, font=F_BASE, fill=None, borde=True, align=None):
    c = ws[ref]
    c.value = valor
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if borde:
        c.border = BORDE
    c.alignment = align or Alignment(vertical='top')
    return c


def encabezado(ws, fila, titulos, col0=1, alto=45):
    for i, t in enumerate(titulos):
        c = ws.cell(row=fila, column=col0 + i, value=t)
        c.font = F_HEADER
        c.fill = FILL_HEADER
        c.border = BORDE
        c.alignment = CENTRO
    ws.row_dimensions[fila].height = alto


def titulo(ws, texto, subtitulo=None):
    ws['A1'] = texto
    ws['A1'].font = F_TITULO
    if subtitulo:
        ws['A2'] = subtitulo
        ws['A2'].font = F_NOTA


def anchos(ws, valores):
    for i, w in enumerate(valores, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def nombre_definido(wb, nombre, hoja, ref):
    wb.defined_names[nombre] = DefinedName(nombre, attr_text=f"'{hoja}'!{ref}")


def hoja_parametros(wb):
    ws = wb.create_sheet('Parámetros')
    titulo(ws, 'Parámetros del cálculo',
           'Celdas en azul claro: valores editables. Todas las hojas del libro los leen por nombre definido.')
    encabezado(ws, 4, ['Parámetro', 'Valor', 'Unidad', 'Nombre definido', 'Fuente / nota'], alto=30)
    filas = [
        ('Dosis de fertilización de nitrógeno', DOSIS_N, 'kg N/ha/año', 'DosisN',
         'Informe Fase 1, Tabla 4.6 (Narváez, 2024)', True, '0.00'),
        ('Dosis de P₂O₅', DOSIS_P2O5, 'kg P₂O₅/ha/año', 'DosisP2O5', 'Informe Fase 1, Tabla 4.6', True, '0.00'),
        ('Factor de conversión P₂O₅ → P', CONV_P2O5_P, 'adimensional', 'ConvP2O5',
         'Informe Fase 1, Tabla 4.6: 2 × masa atómica de P ÷ masa molecular de P₂O₅', True, '0.000'),
        ('Dosis de fósforo elemental', '=DosisP2O5*ConvP2O5', 'kg P/ha/año', 'DosisP',
         'Calculada; el informe reporta 20,056', False, '0.000'),
        ('Fracción de pérdida por escorrentía, N (escenario medio)', FRAC_N, 'fracción', 'FracN',
         'Informe Fase 1, Tabla 4.7 (Scarpare et al., 2023)', True, '0.00%'),
        ('Fracción de pérdida por escorrentía, P (escenario medio)', FRAC_P, 'fracción', 'FracP',
         'Informe Fase 1, Tabla 4.8 (Jeong et al., 2011)', True, '0.00%'),
        ('Factor de emisión N = dosis × fracción', '=DosisN*FracN', 'kg N/ha/año', 'FE_N',
         'Control: 9,568 (nota de la Tabla 4.22)', False, '0.000'),
        ('Factor de emisión P = dosis × fracción', '=DosisP*FracP', 'kg P/ha/año', 'FE_P',
         'Control: 1,376 (nota de la Tabla 4.22)', False, '0.000'),
        None,
        ('Factor de retención: condición base (sin corredor)', RETENCION['base'], 'fracción', 'RetBase',
         'Procedimiento de la Actividad 5', True, '0.00'),
        ('Factor de retención: escenario 1 (corredor joven, remoción 20 %)', RETENCION['esc1'], 'fracción',
         'RetEsc1', 'Procedimiento de la Actividad 5', True, '0.00'),
        ('Factor de retención: escenario 2 (corredor establecido, remoción 40 %)', RETENCION['esc2'],
         'fracción', 'RetEsc2', 'Procedimiento de la Actividad 5', True, '0.00'),
        None,
        ('Días por año', DIAS_ANO, 'd', 'DiasAno', 'Carga anual distribuida de forma uniforme', True, '0'),
        ('Factor de conversión de carga a concentración', CONV_Q, '—', 'ConvQ',
         'C (mg/L) = carga (kg/d) ÷ (Q (m³/s) × 86,4); 86,4 = 86 400 s/d × 1 000 L/m³ ÷ 10⁶ mg/kg',
         True, '0.0'),
        ('Umbral de dominancia de la carga puntual', UMBRAL_DOMINANCIA, 'fracción de la carga total',
         'UmbralDom', 'Se excluye el subtramo si puntual ÷ (puntual + difusa) supera este valor',
         True, '0%'),
        ('Condición hidrológica del caudal de referencia', CONDICION_DEFECTO, 'lista', 'CondHidro',
         'Promedio: misma base temporal que la concentración histórica media. Umbrales: hoja Caudales',
         True, None),
    ]
    fila = 5
    for f in filas:
        if f is None:
            fila += 1
            continue
        nombre, valor, unidad, nd, fuente, editable, fmt = f
        celda(ws, f'A{fila}', nombre, align=WRAP_TOP)
        celda(ws, f'B{fila}', valor, fmt=fmt, fill=FILL_INPUT if editable else None)
        celda(ws, f'C{fila}', unidad, align=WRAP_TOP)
        celda(ws, f'D{fila}', nd)
        celda(ws, f'E{fila}', fuente, align=WRAP_TOP)
        nombre_definido(wb, nd, 'Parámetros', f'$B${fila}')
        if nd == 'CondHidro':
            fila_cond = fila
        fila += 1

    fila += 1
    celda(ws, f'A{fila}', 'Condiciones disponibles', font=F_HEADER, fill=FILL_HEADER)
    ini = fila + 1
    for i, c in enumerate(CONDICIONES):
        celda(ws, f'A{ini + i}', c)
    fin = ini + len(CONDICIONES) - 1
    nombre_definido(wb, 'ListaCond', 'Parámetros', f'$A${ini}:$A${fin}')
    dv = DataValidation(type='list', formula1=f'$A${ini}:$A${fin}', allow_blank=False)
    dv.error = 'Seleccionar una condición de la lista.'
    ws.add_data_validation(dv)
    dv.add(f'B{fila_cond}')
    anchos(ws, [52, 12, 22, 14, 70])
    ws.freeze_panes = 'A5'
    return ws


def hoja_serie(wb, series):
    ws = wb.create_sheet('Serie_diaria')
    encabezado(ws, 1, ['Estación hidrométrica', 'Fecha', 'Caudal (m³/s)'], alto=30)
    fila = 2
    for nombre, serie in series.items():
        for fecha, q in serie:
            ws.cell(row=fila, column=1, value=nombre).font = F_BASE
            c = ws.cell(row=fila, column=2, value=fecha)
            c.font = F_BASE
            c.number_format = 'yyyy-mm-dd'
            c = ws.cell(row=fila, column=3, value=q)
            c.font = F_BASE
            c.number_format = FMT_Q
            fila += 1
    anchos(ws, [24, 12, 14])
    ws.freeze_panes = 'A2'
    return fila - 1   # última fila con datos


def hoja_caudales(wb, estaciones, hidro, ultima):
    ws = wb.create_sheet('Caudales')
    titulo(ws, 'Caudal de referencia por estación hidrométrica y condición hidrológica',
           'Umbrales de la curva de duración (Informe Fase 1, Tabla 3.1): verano Q ≤ umbral de verano; '
           'invierno Q ≥ umbral de invierno; transición entre ambos. Promedios calculados sobre la hoja '
           'Serie_diaria (días con dato válido).')
    encabezado(ws, 4, ['Estación hidrométrica', 'Río', 'Estado', 'Periodo', 'Umbral verano ≤ (m³/s)',
                       'Umbral invierno ≥ (m³/s)', 'n días', 'Q promedio (m³/s)', 'n días verano',
                       'Q verano (m³/s)', 'n días transición', 'Q transición (m³/s)', 'n días invierno',
                       'Q invierno (m³/s)', 'Q de referencia según CondHidro (m³/s)',
                       'Q promedio publicado, Tabla 3.1 (m³/s)', 'Diferencia serie ÷ Tabla 3.1'], alto=60)
    est = f"'Serie_diaria'!$A$2:$A${ultima}"
    val = f"'Serie_diaria'!$C$2:$C${ultima}"
    filas = {}
    for i, nombre in enumerate(estaciones):
        r = 5 + i
        h = hidro[nombre]
        filas[nombre] = r
        celda(ws, f'A{r}', nombre)
        celda(ws, f'B{r}', h['rio'].replace('Rio ', 'Río ') if h['rio'].startswith('Rio ') else 'Río ' + h['rio'])
        celda(ws, f'C{r}', h['estado'])
        celda(ws, f'D{r}', h['años_datos'])
        celda(ws, f'E{r}', h['umbral_verano_m3s'], fmt='0.00', fill=FILL_INPUT)
        celda(ws, f'F{r}', h['umbral_invierno_m3s'], fmt='0.00', fill=FILL_INPUT)
        celda(ws, f'G{r}', f'=COUNTIFS({est},$A{r})', fmt='0')
        celda(ws, f'H{r}', f'=IFERROR(AVERAGEIFS({val},{est},$A{r}),"")', fmt=FMT_Q)
        celda(ws, f'I{r}', f'=COUNTIFS({est},$A{r},{val},"<="&$E{r})', fmt='0')
        celda(ws, f'J{r}', f'=IFERROR(AVERAGEIFS({val},{est},$A{r},{val},"<="&$E{r}),"")', fmt=FMT_Q)
        celda(ws, f'K{r}', f'=COUNTIFS({est},$A{r},{val},">"&$E{r},{val},"<"&$F{r})', fmt='0')
        celda(ws, f'L{r}', f'=IFERROR(AVERAGEIFS({val},{est},$A{r},{val},">"&$E{r},{val},"<"&$F{r}),"")',
              fmt=FMT_Q)
        celda(ws, f'M{r}', f'=COUNTIFS({est},$A{r},{val},">="&$F{r})', fmt='0')
        celda(ws, f'N{r}', f'=IFERROR(AVERAGEIFS({val},{est},$A{r},{val},">="&$F{r}),"")', fmt=FMT_Q)
        celda(ws, f'O{r}', f'=CHOOSE(MATCH(CondHidro,ListaCond,0),H{r},J{r},L{r},N{r})', fmt=FMT_Q)
        celda(ws, f'P{r}', h['promedio_m3s'], fmt='0.00')
        celda(ws, f'Q{r}', f'=IFERROR(H{r}/P{r}-1,"")', fmt=FMT_PCT)
    ultima_fila = 4 + len(estaciones)
    nota = ultima_fila + 2
    ws[f'A{nota}'] = ('La diferencia frente a la Tabla 3.1 proviene del número de días: la tabla publicada cuenta '
                      'también días sin dato válido. El cálculo usa solo días con caudal registrado.')
    ws[f'A{nota}'].font = F_NOTA
    anchos(ws, [18, 16, 11, 11, 11, 11, 8, 11, 8, 11, 9, 11, 8, 11, 14, 13, 12])
    ws.freeze_panes = 'B5'
    return filas, ultima_fila


COLS_MAESTRA = [
    # (clave, título, ancho, formato)
    ('orden', 'N.º', 5, '0'),
    ('rio', 'Río', 13, None),
    ('subtramo', 'Subtramo', 10, '0'),
    ('pos', 'Posición en la Tabla 5.2', 9, None),
    ('delim', 'Estaciones que delimitan el subtramo (inicio → cierre)', 38, None),
    ('cat', 'Categoría de carga de N', 10, None),
    ('ha', 'Área de caña del subtramo (ha)', 11, FMT_HA),
    ('est', 'Estación de calidad seleccionada', 34, None),
    ('km', 'km de la estación sobre el eje', 9, '0.00'),
    ('motivo', 'Motivo de selección', 40, None),
    ('hid', 'Estación hidrométrica asignada', 14, None),
    ('rel', 'Posición de la estación hidrométrica respecto a la de calidad', 30, None),
    ('q', 'Caudal de referencia (m³/s)', 11, FMT_Q),
    ('cn_b', 'Carga N base (kg N/d)', 10, FMT_KGD),
    ('cn_1', 'Carga N escenario 1 (kg N/d)', 10, FMT_KGD),
    ('cn_2', 'Carga N escenario 2 (kg N/d)', 10, FMT_KGD),
    ('cp_b', 'Carga P base (kg P/d)', 10, FMT_KGD),
    ('cp_1', 'Carga P escenario 1 (kg P/d)', 10, FMT_KGD),
    ('cp_2', 'Carga P escenario 2 (kg P/d)', 10, FMT_KGD),
    ('kn_b', 'Concentración N base (mg N/L)', 13, FMT_CONC),
    ('kn_1', 'Concentración N escenario 1 (mg N/L)', 13, FMT_CONC),
    ('kn_2', 'Concentración N escenario 2 (mg N/L)', 13, FMT_CONC),
    ('kp_b', 'Concentración P base (mg P/L)', 13, FMT_CONC),
    ('kp_1', 'Concentración P escenario 1 (mg P/L)', 13, FMT_CONC),
    ('kp_2', 'Concentración P escenario 2 (mg P/L)', 13, FMT_CONC),
    ('nt', 'N total histórico, media (mg N/L)', 11, FMT_CONC),
    ('n_nt', 'n registros NT', 9, '0'),
    ('pt', 'P total histórico, media (mg P/L)', 11, FMT_CONC),
    ('n_pt', 'n registros PT', 9, '0'),
    ('rat_n', 'Contraste N: calculada base ÷ histórica', 11, FMT_RAT),
    ('rat_p', 'Contraste P: calculada base ÷ histórica', 11, FMT_RAT),
    ('dom', 'Carga puntual ÷ carga total, N', 11, FMT_PCT),
    ('crit', 'Criterio de dominancia', 34, None),
    ('dom_sn', 'Dominancia puntual (> umbral)', 11, None),
    ('estado', 'Estado en el monitoreo', 11, None),
    ('mot_exc', 'Motivo de exclusión / observación', 44, None),
]
COL = {c[0]: get_column_letter(i + 1) for i, c in enumerate(COLS_MAESTRA)}
FILA0_MAESTRA = 5


def hoja_maestra(wb, ejercicio, filas_caudal, ultima_caudal, filas_dom):
    ws = wb.create_sheet('Tabla_maestra', 0)
    titulo(ws, 'Tabla maestra: traducción de carga difusa a concentración por subtramo (Actividad 5)',
           'Subtramos en categoría de carga de N Alta o Muy alta, excepto los ríos con exclusión previa '
           '(hoja Exclusiones_previas). Carga = ha × dosis × fracción de pérdida × factor de retención ÷ 365; '
           'concentración = carga ÷ (Q × 86,4).')
    encabezado(ws, 4, [c[1] for c in COLS_MAESTRA], alto=75)
    cau_a = f"'Caudales'!$A$5:$A${ultima_caudal}"
    cau_o = f"'Caudales'!$O$5:$O${ultima_caudal}"
    ret = {'b': 'RetBase', '1': 'RetEsc1', '2': 'RetEsc2'}
    for i, e in enumerate(ejercicio):
        r = FILA0_MAESTRA + i
        e['fila_excel'] = r
        C = {k: f'{v}{r}' for k, v in COL.items()}
        v = {
            'orden': i + 1, 'rio': e['rio'], 'subtramo': e['subtramo'],
            'pos': int(e['posicion_tabla_5_2']) if e['posicion_tabla_5_2'].isdigit() else e['posicion_tabla_5_2'],
            'delim': e['delimitacion'], 'cat': e['categoria'], 'ha': e['area_cana_ha'],
            'est': e['estacion'], 'km': e['km_estacion'], 'motivo': e['motivo_seleccion'],
            'hid': e['estacion_hidro'], 'rel': e['relacion_hidro'],
            'q': f'=INDEX({cau_o},MATCH({C["hid"]},{cau_a},0))',
            'nt': e['nt'], 'n_nt': e['n_nt'], 'pt': e['pt'], 'n_pt': e['n_pt'],
        }
        for s in ('b', '1', '2'):
            v[f'cn_{s}'] = f'={C["ha"]}*DosisN*FracN*{ret[s]}/DiasAno'
            v[f'cp_{s}'] = f'={C["ha"]}*DosisP*FracP*{ret[s]}/DiasAno'
            v[f'kn_{s}'] = f'=IF(ISNUMBER({C["q"]}),{C[f"cn_{s}"]}/({C["q"]}*ConvQ),"")'
            v[f'kp_{s}'] = f'=IF(ISNUMBER({C["q"]}),{C[f"cp_{s}"]}/({C["q"]}*ConvQ),"")'
        v['rat_n'] = f'=IF(AND(ISNUMBER({C["kn_b"]}),N({C["nt"]})>0),{C["kn_b"]}/{C["nt"]},"")'
        v['rat_p'] = f'=IF(AND(ISNUMBER({C["kp_b"]}),N({C["pt"]})>0),{C["kp_b"]}/{C["pt"]},"")'
        ref_frac, ref_crit = filas_dom[(e['clave'], e['subtramo'])]
        v['dom'] = f"='Carga_puntual'!{ref_frac}"
        v['crit'] = f"='Carga_puntual'!{ref_crit}"
        v['dom_sn'] = f'=IF(ISNUMBER({C["dom"]}),IF({C["dom"]}>UmbralDom,"Sí","No"),"Pendiente")'
        v['estado'] = (f'=IF({C["dom_sn"]}="Pendiente","Pendiente",'
                       f'IF({C["dom_sn"]}="Sí","Excluida","Incluida"))')
        v['mot_exc'] = (f'=IF({C["estado"]}="Excluida","Carga puntual dominante ("&ROUND({C["dom"]}*100,1)'
                        f'&" % de la carga total de N; "&{C["crit"]}&")",'
                        f'IF({C["estado"]}="Pendiente",{C["crit"]},""))')
        for k, (_, _, _, fmt) in zip(COL, COLS_MAESTRA):
            celda(ws, C[k], v[k], fmt=fmt, align=WRAP_TOP)
    anchos(ws, [c[2] for c in COLS_MAESTRA])
    ws.freeze_panes = f'{COL["est"]}{FILA0_MAESTRA}'
    return ws


COLS_VERT = [
    # (clave, título, ancho, formato)
    ('registro', 'Registro', 9, None),
    ('cuenca', 'Cuenca (inventario)', 13, None),
    ('tipo', 'Tipo de vertimiento', 12, None),
    ('receptor', 'Cuerpo de agua receptor', 22, None),
    ('tramo_inv', 'Tramo (inventario)', 9, None),
    ('estacion_inv', 'Estación aguas abajo del vertimiento (inventario)', 40, None),
    ('estacion_eq', 'Estación equivalente de la red de calidad', 36, None),
    ('km', 'km de la estación sobre el eje', 9, '0.00'),
    ('subtramo', 'Subtramo asignado (criterio por estación)', 16, None),
    ('control', 'Subtramo según coordenadas (control)', 18, None),
    ('q_l_s', 'Caudal sujeto al cobro (L/s)', 11, '#,##0.00'),
    ('dbo_kg_ano', 'Carga DBO5 (kg/año)', 13, '#,##0'),
    ('sst_kg_ano', 'Carga SST (kg/año)', 13, '#,##0'),
    ('dbo_kg_d', 'Carga DBO5 (kg/d)', 11, FMT_KGD),
    ('sst_kg_d', 'Carga SST (kg/d)', 11, FMT_KGD),
    ('observacion', 'Observación del control espacial', 50, None),
]
COLV = {c[0]: get_column_letter(i + 1) for i, c in enumerate(COLS_VERT)}
FILA0_VERT = 5


def hoja_vertimientos(wb, listado):
    ws = wb.create_sheet('Vertimientos')
    titulo(ws, 'Inventario de tasas retributivas (CVC): vertimientos de las cuencas de los subtramos pendientes',
           'Fuente: docs/Tasas retributivas.xlsx. Asignación: la estación aguas abajo del vertimiento, llevada a la '
           'red de calidad (data/fuentes/monitoreo_corredor/equivalencia_estaciones_tasas.csv), debe quedar entre la '
           'estación de inicio (exclusiva) y la de cierre (inclusiva) del subtramo. Las coordenadas (EPSG:9377) solo '
           'se usan como control y no cambian la asignación.')
    encabezado(ws, 4, [c[1] for c in COLS_VERT], alto=60)
    orden = sorted(listado, key=lambda v: (v['subtramo'] == '', v['subtramo'], v['cuenca'], v['registro']))
    for i, v in enumerate(orden):
        r = FILA0_VERT + i
        vals = {k: v.get(k) for k in COLV}
        vals['km'] = v['km'] if v['km'] is not None else ''
        vals['subtramo'] = v['subtramo'] or '—'
        vals['control'] = v['control'] or 'Fuera de la franja de 700 m'
        vals['dbo_kg_d'] = f'={COLV["dbo_kg_ano"]}{r}/DiasAno'
        vals['sst_kg_d'] = f'={COLV["sst_kg_ano"]}{r}/DiasAno'
        for k, (_, _, _, fmt) in zip(COLV, COLS_VERT):
            celda(ws, f'{COLV[k]}{r}', vals[k], fmt=fmt, align=WRAP_TOP)
    anchos(ws, [c[2] for c in COLS_VERT])
    ws.freeze_panes = f'B{FILA0_VERT}'
    return FILA0_VERT + len(orden) - 1


def hoja_carga_puntual(wb, ejercicio, tabla_4_10, ultima_vert):
    ws = wb.create_sheet('Carga_puntual')
    titulo(ws, 'Dominancia de la carga puntual por subtramo',
           'Bloque 1: Tabla 4.10 del Informe Fase 1 (PORH Bolo y Fraile, modelo QUAL2Kw), digitalizada. '
           'Una carga difusa de 0 indica un segmento no definido en el archivo de modelación, no ausencia de aporte '
           '(Informe Fase 1, p. 166).')
    encabezado(ws, 4, ['Nutriente', 'Campaña', 'Río', 'Tramo PORH', 'Carga vertimientos (kg/d)',
                       'Carga afluentes (kg/d)', 'Carga puntual, vert. + afl. (kg/d)',
                       'Carga difusa QUAL2Kw (kg/d)', 'Difusa definida', 'Puntual ÷ (puntual + difusa)',
                       'Subtramo del geovisor', 'Correspondencia de la delimitación'], alto=60)
    fila_de = {}
    for i, t in enumerate(tabla_4_10):
        r = 5 + i
        fila_de[(t['nutriente'], t['campana'], normalizar_rio(t['rio_geovisor']), int(t['subtramo_geovisor']))] = r
        celda(ws, f'A{r}', t['nutriente'])
        celda(ws, f'B{r}', t['campana'])
        celda(ws, f'C{r}', t['rio'])
        celda(ws, f'D{r}', t['tramo_porh'])
        for col, k in zip('EFGH', ('carga_vertimientos_kg_d', 'carga_afluentes_kg_d',
                                   'carga_puntual_kg_d', 'carga_difusa_qual2kw_kg_d')):
            celda(ws, f'{col}{r}', float(t[k]), fmt=FMT_KGD, fill=FILL_INPUT)
        celda(ws, f'I{r}', f'=IF(H{r}>0,"Sí","No")')
        celda(ws, f'J{r}', f'=IF(H{r}>0,G{r}/(G{r}+H{r}),"")', fmt=FMT_PCT)
        celda(ws, f'K{r}', f"{t['rio_geovisor']} {t['subtramo_geovisor']}")
        celda(ws, f'L{r}', t['correspondencia'])
    r = 5 + len(tabla_4_10) + 1
    ws[f'A{r}'] = ('Correspondencia parcial: la estación GF2 del PORH es La Industria, mientras que el subtramo 2 '
                   'del geovisor arranca en Puente Vía a Miranda; GF1 → GF2 y GF2 → GF3 se asocian a los '
                   'subtramos 1 y 2 del Fraile de manera aproximada.')
    ws[f'A{r}'].font = F_NOTA

    # Bloque 2: síntesis por subtramo con Tabla 4.10
    r += 2
    ws[f'A{r}'] = ('Bloque 2: síntesis por subtramo, N total (criterio de exclusión) y P total (apoyo). Regla: '
                   'máximo entre campañas con difusa definida; si ninguna la tiene, contraste con la carga difusa '
                   'por factor de emisión (criterio aproximado).')
    ws[f'A{r}'].font = F_BOLD
    r += 1
    encabezado(ws, r, ['Subtramo', 'Tramo PORH', 'Correspondencia', 'N: fracción puntual C1 (difusa QUAL2Kw)',
                       'N: fracción puntual C2 (difusa QUAL2Kw)', 'Carga difusa por factor de emisión, N (kg/d)',
                       'N: fracción puntual C1 frente a FE', 'N: fracción puntual C2 frente a FE',
                       'P: fracción puntual C1 (difusa QUAL2Kw)', 'P: fracción puntual C2 (difusa QUAL2Kw)',
                       'P: fracción adoptada (apoyo)', 'N: fracción puntual adoptada', 'Criterio'], alto=75)
    filas_dom = {}   # (clave, subtramo) → (celda de fracción adoptada, celda de criterio)
    con_porh = [e for e in ejercicio if e.get('tramo_porh')]
    for e in con_porh:
        r += 1
        k = (e['clave'], e['subtramo'])
        f = lambda nut, camp: fila_de[(nut, camp, *k)]
        n1, n2 = f('Nitrógeno total', 'C1 seca'), f('Nitrógeno total', 'C2 húmeda')
        p1, p2 = f('Fósforo total', 'C1 seca'), f('Fósforo total', 'C2 húmeda')
        fe_n = f"'Tabla_maestra'!{COL['cn_b']}{e['fila_excel']}"
        fe_p = f"'Tabla_maestra'!{COL['cp_b']}{e['fila_excel']}"
        celda(ws, f'A{r}', f"{e['rio']} {e['subtramo']}")
        celda(ws, f'B{r}', e['tramo_porh'])
        celda(ws, f'C{r}', e['correspondencia_porh'])
        celda(ws, f'D{r}', f'=J{n1}', fmt=FMT_PCT)
        celda(ws, f'E{r}', f'=J{n2}', fmt=FMT_PCT)
        celda(ws, f'F{r}', f'={fe_n}', fmt=FMT_KGD)
        celda(ws, f'G{r}', f'=G{n1}/(G{n1}+F{r})', fmt=FMT_PCT)
        celda(ws, f'H{r}', f'=G{n2}/(G{n2}+F{r})', fmt=FMT_PCT)
        celda(ws, f'I{r}', f'=J{p1}', fmt=FMT_PCT)
        celda(ws, f'J{r}', f'=J{p2}', fmt=FMT_PCT)
        celda(ws, f'K{r}', f'=IF(COUNT(I{r}:J{r})>0,MAX(I{r}:J{r}),'
                           f'MAX(G{p1}/(G{p1}+{fe_p}),G{p2}/(G{p2}+{fe_p})))', fmt=FMT_PCT)
        celda(ws, f'L{r}', f'=IF(COUNT(D{r}:E{r})>0,MAX(D{r}:E{r}),MAX(G{r}:H{r}))', fmt=FMT_PCT)
        celda(ws, f'M{r}',
              f'=IF(COUNT(D{r}:E{r})>0,"Riguroso: Tabla 4.10 (puntual frente a difusa QUAL2Kw)",'
              f'"Aproximado: difusa QUAL2Kw no definida en el tramo; contraste con la difusa por factor de emisión")',
              align=WRAP_TOP)
        filas_dom[k] = (f'$L${r}', f'$M${r}')

    # Bloque 3: ríos sin balance de masas — inventario de tasas retributivas
    r += 3
    ws[f'A{r}'] = ('Bloque 3: ríos sin balance de masas. Carga puntual del inventario de tasas retributivas de la '
                   'CVC (hoja Vertimientos), expresada en DBO5 y SST. El inventario no trae N ni P y el proyecto no '
                   'tiene un factor validado de conversión DBO5 → N, de modo que la dominancia no se calcula: la '
                   'carga N difusa y la carga DBO5 puntual se presentan lado a lado para su lectura directa.')
    ws[f'A{r}'].font = F_BOLD
    ws[f'A{r}'].alignment = Alignment(wrap_text=False)
    r += 1
    encabezado(ws, r, ['Subtramo', 'N.º de vertimientos asignados', 'Caudal vertido total (L/s)',
                       'Carga DBO5 puntual (kg/d)', 'Carga SST puntual (kg/d)',
                       'Carga N difusa por factor de emisión (kg/d)',
                       'Carga puntual de N (kg/d): solo con un factor validado',
                       'N: fracción puntual adoptada', 'Criterio de dominancia'], alto=75)
    vs = lambda clave: f"'Vertimientos'!${COLV[clave]}${FILA0_VERT}:${COLV[clave]}${ultima_vert}"
    pendientes = [e for e in ejercicio if not e.get('tramo_porh')]
    for e in pendientes:
        r += 1
        k = (e['clave'], e['subtramo'])
        celda(ws, f'A{r}', e['etiqueta'])
        celda(ws, f'B{r}', f'=COUNTIFS({vs("subtramo")},A{r})', fmt='0')
        celda(ws, f'C{r}', f'=SUMIFS({vs("q_l_s")},{vs("subtramo")},A{r})', fmt='#,##0.00')
        celda(ws, f'D{r}', f'=SUMIFS({vs("dbo_kg_d")},{vs("subtramo")},A{r})', fmt=FMT_KGD)
        celda(ws, f'E{r}', f'=SUMIFS({vs("sst_kg_d")},{vs("subtramo")},A{r})', fmt=FMT_KGD)
        celda(ws, f'F{r}', f"='Tabla_maestra'!{COL['cn_b']}{e['fila_excel']}", fmt=FMT_KGD)
        celda(ws, f'G{r}', None, fmt=FMT_KGD, fill=FILL_INPUT)
        celda(ws, f'H{r}', f'=IF(ISNUMBER(G{r}),G{r}/(G{r}+F{r}),"")', fmt=FMT_PCT)
        celda(ws, f'I{r}', f'=IF(ISNUMBER(H{r}),"{CRITERIO_N_MANUAL}",'
                           f'IF(B{r}=0,"{CRITERIO_SIN_VERTIMIENTOS}","{CRITERIO_CON_VERTIMIENTOS}"))',
              align=WRAP_TOP)
        filas_dom[k] = (f'$H${r}', f'$I${r}')
        e['refs_b3'] = {'n_vertimientos': f"'Carga_puntual'!B{r}", 'q_vertido_l_s': f"'Carga_puntual'!C{r}",
                        'carga_dbo5_kg_d': f"'Carga_puntual'!D{r}", 'carga_sst_kg_d': f"'Carga_puntual'!E{r}"}
    r += 1
    ws[f'A{r + 1}'] = ('La columna G queda vacía a propósito: solo debe diligenciarse con una carga de N obtenida con '
                       'un factor de conversión validado por la dirección del proyecto. Al diligenciarla, la fracción '
                       'y el estado del subtramo se recalculan en la tabla maestra.')
    ws[f'A{r + 1}'].font = F_NOTA
    r += 2
    for e in pendientes:
        for nota in e.get('notas_vertimientos', []):
            r += 1
            ws[f'A{r}'] = f"{e['etiqueta']}. {nota}"
            ws[f'A{r}'].font = F_NOTA
    anchos(ws, [18, 14, 13, 13, 13, 13, 13, 13, 34, 13, 13, 13, 40])
    ws.freeze_panes = 'A5'
    return filas_dom


def hoja_exclusiones(wb, previas):
    ws = wb.create_sheet('Exclusiones_previas')
    titulo(ws, 'Subtramos Alta / Muy alta de ríos excluidos antes del cálculo',
           'Exclusión por río completo, previa a la selección de estaciones (procedimiento de la Actividad 5).')
    encabezado(ws, 4, ['Río', 'Subtramo', 'Posición en la Tabla 5.2', 'Categoría de carga de N',
                       'Área de caña (ha)', 'Carga N base (kg N/d)', 'Estación de cierre', 'Motivo'], alto=45)
    for i, e in enumerate(previas):
        r = 5 + i
        celda(ws, f'A{r}', e['rio'])
        celda(ws, f'B{r}', e['subtramo'])
        celda(ws, f'C{r}', e['posicion_tabla_5_2'])
        celda(ws, f'D{r}', e['categoria'])
        celda(ws, f'E{r}', e['area_cana_ha'], fmt=FMT_HA)
        celda(ws, f'F{r}', f'=E{r}*DosisN*FracN*RetBase/DiasAno', fmt=FMT_KGD)
        celda(ws, f'G{r}', e['estacion'], align=WRAP_TOP)
        celda(ws, f'H{r}', e['motivo_exclusion'], align=WRAP_TOP)
    ws[f'A{6 + len(previas)}'] = ('Los datos de caudal del PORH / QUAL2Kw (usados solo para Bolo y Fraile en la '
                                   'Tabla 4.10) son una fuente distinta y no se emplean para suplir la falta de '
                                   'caudal de estación.')
    ws[f'A{6 + len(previas)}'].font = F_NOTA
    anchos(ws, [14, 9, 10, 11, 11, 11, 40, 55])


def hoja_verificacion(wb, ejercicio):
    ws = wb.create_sheet('Verificación')
    titulo(ws, 'Verificación: fórmulas del libro frente al cálculo independiente en Python',
           f'Válida con los parámetros por defecto (condición {CONDICION_DEFECTO}, retención 1 / 0,8 / 0,6). '
           'Si se editan parámetros, las diferencias dejan de ser cero por diseño.')
    encabezado(ws, 4, ['Subtramo', 'Magnitud', 'Valor en el libro', 'Valor calculado en Python',
                       'Diferencia absoluta'], alto=30)
    magnitudes = [
        ('q', 'Caudal de referencia (m³/s)', 'q_ref'),
        ('cn_b', 'Carga N base (kg/d)', 'carga_n_base'),
        ('cn_1', 'Carga N escenario 1 (kg/d)', 'carga_n_esc1'),
        ('cn_2', 'Carga N escenario 2 (kg/d)', 'carga_n_esc2'),
        ('cp_b', 'Carga P base (kg/d)', 'carga_p_base'),
        ('kn_b', 'Concentración N base (mg/L)', 'conc_n_base'),
        ('kn_1', 'Concentración N escenario 1 (mg/L)', 'conc_n_esc1'),
        ('kn_2', 'Concentración N escenario 2 (mg/L)', 'conc_n_esc2'),
        ('kp_b', 'Concentración P base (mg/L)', 'conc_p_base'),
        ('kp_2', 'Concentración P escenario 2 (mg/L)', 'conc_p_esc2'),
        ('dom', 'Carga puntual ÷ total, N', 'frac_puntual_n'),
        ('crit', 'Criterio de dominancia', 'criterio_dominancia'),
        ('estado', 'Estado', 'estado'),
    ]
    extras = [
        ('n_vertimientos', 'N.º de vertimientos asignados'),
        ('q_vertido_l_s', 'Caudal vertido total (L/s)'),
        ('carga_dbo5_kg_d', 'Carga DBO5 puntual (kg/d)'),
        ('carga_sst_kg_d', 'Carga SST puntual (kg/d)'),
    ]
    r = 4
    for e in ejercicio:
        for k, nombre, clave in magnitudes:
            r += 1
            celda(ws, f'A{r}', f"{e['rio']} {e['subtramo']}")
            celda(ws, f'B{r}', nombre)
            celda(ws, f'C{r}', f"='Tabla_maestra'!{COL[k]}{e['fila_excel']}", fmt='0.000000')
            py = e[clave]
            celda(ws, f'D{r}', py if py is not None else '', fmt='0.000000')
            celda(ws, f'E{r}', f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),ABS(C{r}-D{r}),IF(C{r}&""=D{r}&"",0,1))',
                  fmt='0.000000000')
        for clave, nombre in extras if e.get('refs_b3') else []:
            r += 1
            celda(ws, f'A{r}', f"{e['rio']} {e['subtramo']}")
            celda(ws, f'B{r}', nombre)
            celda(ws, f'C{r}', f"={e['refs_b3'][clave]}", fmt='0.000000')
            celda(ws, f'D{r}', e[clave], fmt='0.000000')
            celda(ws, f'E{r}', f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),ABS(C{r}-D{r}),IF(C{r}&""=D{r}&"",0,1))',
                  fmt='0.000000000')
    celda(ws, 'G4', 'Máxima diferencia', font=F_HEADER, fill=FILL_HEADER)
    celda(ws, 'H4', f'=MAX(E5:E{r})', fmt='0.000000000')
    for row in ws.iter_rows(min_row=5, max_row=r, min_col=3, max_col=4):
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    anchos(ws, [18, 34, 22, 22, 16, 2, 18, 14])
    ws.freeze_panes = 'A5'


def hoja_leame(wb, fe_n, fe_p, n_ejercicio, n_previas, compartidas):
    ws = wb.create_sheet('Léame', 0)
    ws['A1'] = 'Traducción de carga difusa a concentración: monitoreo del corredor biológico (Actividad 5)'
    ws['A1'].font = F_TITULO
    ws['A2'] = (f'Proyecto 890K, UAO × ASOCAÑA. Generado el {datetime.date.today().isoformat()} con '
                'scripts/build_monitoreo_corredor.py (repositorio Rio_Cauca_Baseline).')
    ws['A2'].font = F_NOTA
    secciones = [
        ('Objetivo', [
            'Traducir la carga difusa de nitrógeno y fósforo total atribuible a la caña de azúcar de cada subtramo '
            'priorizado en una concentración esperada en su estación de calidad del agua, para la condición sin '
            'corredor y para dos escenarios de retención del corredor biológico, y contrastarla con la '
            'concentración histórica observada.',
        ]),
        ('Universo', [
            f'Subtramos en categoría de carga de N Alta o Muy alta de la Tabla 5.2 (con la caña actualizada en '
            f'septiembre de 2026): {n_ejercicio + n_previas} de 46.',
            'Exclusión previa de ríos completos: Nima, Guachal y Zabaletas, por ausencia de datos de caudal en las '
            'fuentes de estación; Risaralda, por ausencia de monitoreo de nitrógeno y fósforo total '
            f'({n_previas} subtramos, hoja Exclusiones_previas). Quedan {n_ejercicio} subtramos en el ejercicio.',
        ]),
        ('Selección de estaciones', [
            'Estación de calidad: la estación de cierre del subtramo, que es la más aguas abajo del tramo, con '
            'registros de nitrógeno total y fósforo total. Los subtramos están delimitados por estaciones, de modo '
            'que la correspondencia es uno a uno.'
            + (' Estaciones compartidas: ' + '; '.join(f'{k}: {", ".join(v)}' for k, v in compartidas.items())
               if compartidas else ' Ninguna estación de calidad queda compartida entre subtramos.'),
            'Estación hidrométrica: la más cercana, sobre el eje del mismo río, a la estación de calidad, entre las '
            'que tienen al menos cinco años de serie diaria. La posición relativa se declara en la tabla maestra. '
            'Cuando la estación hidrométrica queda aguas arriba, el caudal se subestima y la concentración calculada '
            'resulta mayor que la que correspondería al punto de calidad.',
        ]),
        ('Fórmulas', [
            'Carga (kg/d) = área de caña del subtramo (ha) × dosis de fertilización (kg/ha/año) × fracción de pérdida '
            'por escorrentía × factor de retención del corredor ÷ 365.',
            f'El producto dosis × fracción es el factor de emisión del escenario medio del Capítulo 4: '
            f'{fe_n:.3f} kg N/ha/año (184 × 5,20 %) y {fe_p:.3f} kg P/ha/año (46 × 0,436 × 6,86 %). La fracción de '
            'pérdida se aplica una sola vez; el factor de retención es adicional y no la reemplaza.',
            'Factor de retención: 1 (condición base, sin corredor), 0,8 (escenario 1, corredor joven, remoción del '
            '20 %) y 0,6 (escenario 2, corredor establecido, remoción del 40 %).',
            'Concentración (mg/L) = carga (kg/d) ÷ (caudal de referencia (m³/s) × 86,4).',
            'Caudal de referencia: media de la serie diaria de la estación hidrométrica para la condición elegida en '
            'Parámetros (Promedio por defecto; Verano, Transición o Invierno según los umbrales de la Tabla 3.1).',
        ]),
        ('Contraste con el histórico', [
            'Concentración histórica: media de todas las campañas de la estación, con los valores bajo el límite de '
            'detección incorporados con el valor del propio límite (mismo criterio que el Capítulo 3 y que '
            'data/calidad_agua/estadisticas_puntos.csv). El Capítulo 3 reporta nitratos (N-NO₃); aquí se usa '
            'nitrógeno total, que tiene menos registros por estación (columna n).',
            'Se espera que la concentración calculada sea menor que la histórica, porque el histórico integra carga '
            'puntual y otros aportes difusos que este cálculo no captura. El cociente se reporta como contraste, no '
            'como error.',
            'La concentración calculada corresponde al aporte de la caña del propio subtramo; no acumula el de los '
            'subtramos situados aguas arriba.',
        ]),
        ('Dominancia de la carga puntual', [
            'Umbral: se excluye del monitoreo el subtramo cuya carga puntual supera el 50 % de la carga total '
            '(puntual + difusa) de nitrógeno total. El fósforo total se reporta como apoyo.',
            'Bolo y Fraile (criterio riguroso): Tabla 4.10 del Informe Fase 1, balance de masas con el modelo '
            'QUAL2Kw del PORH. Se toma el máximo entre las campañas con carga difusa definida. Si el modelo no '
            'define la difusa del tramo en ninguna campaña, la carga puntual se contrasta con la carga difusa por '
            'factor de emisión y el caso se marca como criterio aproximado.',
            'Resto de ríos: inventario de tasas retributivas de la CVC (hoja Vertimientos). Cada vertimiento se '
            'asigna al subtramo cuya delimitación (estación de inicio exclusiva, estación de cierre inclusiva) '
            'contiene su estación aguas abajo, llevada a la red de calidad mediante una tabla de equivalencias de '
            'nombres. Las coordenadas del inventario se usan solo como control espacial.',
            'El inventario reporta DBO5 y SST, no nitrógeno ni fósforo, y el proyecto no dispone de un factor '
            'validado de conversión. Por eso no se calcula la dominancia en estos subtramos: el bloque 3 de la hoja '
            'Carga_puntual presenta lado a lado la carga N difusa y la carga DBO5 puntual, y el estado queda '
            'Pendiente. Se distinguen dos situaciones: subtramos con vertimientos identificados (pendientes de un '
            'factor DBO5 → N) y subtramos sin vertimientos puntuales digitalizados en el inventario, que no equivale '
            'a ausencia de fuentes puntuales.',
        ]),
        ('Estructura del libro', [
            'Tabla_maestra: resultado por subtramo. Parámetros: valores editables. Caudales y Serie_diaria: caudal '
            'de referencia por condición. Carga_puntual: Tabla 4.10 y síntesis del inventario de tasas retributivas. '
            'Vertimientos: registros del inventario y su asignación. Exclusiones_previas: ríos fuera del ejercicio. '
            'Verificación: fórmulas frente a cálculo independiente.',
            'Las celdas en azul claro son editables; el resto son fórmulas o datos de origen. Cambiar un parámetro, '
            'un umbral de caudal o la condición hidrológica recalcula el libro en cascada.',
        ]),
        ('Fuentes', [
            'Informe Fase 1 v1 (2026-08-28): Tablas 3.1, 4.1, 4.6, 4.7, 4.8, 4.10, 4.22 y 5.2.',
            'Repositorio Rio_Cauca_Baseline: docs/tramos_cana_tributarios.csv (caña por tramo), '
            'data/fuentes/priorizacion/Priorizacion_NP_subtramos.csv (categorías), data/hidrologia/ (caudal diario '
            'CVC), data/calidad_agua/csv_por_punto/ (calidad del agua), '
            'data/fuentes/monitoreo_corredor/ (posiciones sobre el eje, Tabla 4.10 digitalizada y equivalencia de '
            'estaciones del inventario), docs/Tasas retributivas.xlsx (inventario CVC de vertimientos).',
        ]),
    ]
    r = 4
    for tit, parrafos in secciones:
        ws[f'A{r}'] = tit
        ws[f'A{r}'].font = F_BOLD
        r += 1
        for p in parrafos:
            ws[f'A{r}'] = p
            ws[f'A{r}'].font = F_BASE
            ws[f'A{r}'].alignment = WRAP_TOP
            ws.row_dimensions[r].height = 15 * max(1, math.ceil(len(p) / 125))
            r += 1
        r += 1
    ws.column_dimensions['A'].width = 130


def escribir_excel(ejercicio, previas, compartidas, tabla_4_10, series, hidro, fe_n, fe_p, listado_vert):
    wb = Workbook()
    wb.remove(wb.active)
    hoja_parametros(wb)
    ultima_serie = hoja_serie(wb, series)
    filas_caudal, ultima_caudal = hoja_caudales(wb, list(series), hidro, ultima_serie)
    # La tabla maestra necesita las filas de Carga_puntual y viceversa: se
    # asignan primero las filas de la maestra (orden fijo) y luego se escribe.
    for i, e in enumerate(ejercicio):
        e['fila_excel'] = FILA0_MAESTRA + i
    ultima_vert = hoja_vertimientos(wb, listado_vert)
    filas_dom = hoja_carga_puntual(wb, ejercicio, tabla_4_10, ultima_vert)
    hoja_maestra(wb, ejercicio, filas_caudal, ultima_caudal, filas_dom)
    hoja_exclusiones(wb, previas)
    hoja_verificacion(wb, ejercicio)
    hoja_leame(wb, fe_n, fe_p, len(ejercicio), len(previas), compartidas)
    orden = ['Léame', 'Tabla_maestra', 'Parámetros', 'Carga_puntual', 'Vertimientos', 'Caudales',
             'Exclusiones_previas', 'Verificación', 'Serie_diaria']
    wb._sheets = [wb[n] for n in orden]
    wb.active = 1
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    wb.save(XLSX_OUT)


def recalcular_con_excel(path):
    """Recalcula y guarda con Excel (COM) para que el archivo lleve valores en
    caché: openpyxl solo escribe fórmulas, y en Vista protegida Excel no
    recalcula, así que las celdas se verían vacías. Sin Excel/pywin32 se omite
    y el libro se recalcula igual al abrirlo (fullCalcOnLoad)."""
    try:
        import win32com.client
    except ImportError:
        return 'omitido (pywin32 no disponible)'
    xl = None
    try:
        xl = win32com.client.DispatchEx('Excel.Application')
        xl.Visible = False
        xl.DisplayAlerts = False
        wb = xl.Workbooks.Open(os.path.abspath(path))
        xl.CalculateFull()
        maxdif = wb.Worksheets('Verificación').Range('H4').Value
        wb.Save()
        wb.Close(False)
        return f'recalculado con Excel; máxima diferencia fórmulas vs. Python = {maxdif:.2e}'
    except Exception as exc:   # Excel no instalado u otro fallo de COM
        return f'omitido ({exc.__class__.__name__}: {exc})'
    finally:
        if xl is not None:
            xl.Quit()


# ── GeoJSON ────────────────────────────────────────────────────────────
def redondear(x, d):
    return None if x is None else round(x, d)


def escribir_geojson(ejercicio, previas, puntos, fe_n, fe_p):
    features = []
    for e in ejercicio:
        props = {
            'grupo': 'ejercicio',
            'rio': e['rio'], 'rio_display': rio_display(e['rio']),
            'etiqueta': f"{rio_display(e['rio'])[4:]} {e['subtramo']}",
            'subtramo': e['subtramo'], 'posicion_tabla_5_2': e['posicion_tabla_5_2'],
            'categoria': e['categoria'], 'delimitacion': e['delimitacion'],
            'area_cana_ha': round(e['area_cana_ha'], 2),
            'estacion': e['estacion'], 'km_estacion': e['km_estacion'],
            'motivo_seleccion': e['motivo_seleccion'],
            'estacion_hidro': e['estacion_hidro'], 'relacion_hidro': e['relacion_hidro'],
            # Media de la serie diaria por condición (hoja Caudales), sin
            # redondear: el panel del geovisor recalcula con estos valores.
            'q_promedio_m3s': e['q_cond']['Promedio'],
            'q_verano_m3s': e['q_cond']['Verano'],
            'q_transicion_m3s': e['q_cond']['Transición'],
            'q_invierno_m3s': e['q_cond']['Invierno'],
        }
        for esc in RETENCION:
            props[f'carga_n_{esc}_kg_d'] = redondear(e[f'carga_n_{esc}'], 2)
        for esc in RETENCION:
            props[f'carga_p_{esc}_kg_d'] = redondear(e[f'carga_p_{esc}'], 2)
        for esc in RETENCION:
            props[f'conc_n_{esc}_mg_l'] = redondear(e[f'conc_n_{esc}'], 4)
        for esc in RETENCION:
            props[f'conc_p_{esc}_mg_l'] = redondear(e[f'conc_p_{esc}'], 4)
        props.update({
            'nt_hist_mg_l': e['nt'], 'n_nt': e['n_nt'], 'pt_hist_mg_l': e['pt'], 'n_pt': e['n_pt'],
            'cociente_n': redondear(e['cociente_n'], 3), 'cociente_p': redondear(e['cociente_p'], 3),
            'pct_puntual_n': redondear(e['frac_puntual_n'] * 100 if e['frac_puntual_n'] is not None else None, 1),
            'pct_puntual_p': redondear(e['frac_puntual_p'] * 100 if e['frac_puntual_p'] is not None else None, 1),
            'tramo_porh': e['tramo_porh'] or None,
            'situacion_puntual': e.get('situacion_puntual'),
            'n_vertimientos': e.get('n_vertimientos'),
            'registros_vertimientos': e.get('registros_vertimientos'),
            'q_vertido_l_s': redondear(e.get('q_vertido_l_s'), 2),
            'carga_dbo5_kg_d': redondear(e.get('carga_dbo5_kg_d'), 2),
            'carga_sst_kg_d': redondear(e.get('carga_sst_kg_d'), 2),
            'criterio_dominancia': e['criterio_dominancia'],
            'estado': e['estado'], 'motivo_exclusion': e['motivo_exclusion'] or None,
        })
        features.append({'type': 'Feature',
                         'geometry': {'type': 'Point', 'coordinates': [round(e['lon'], 6), round(e['lat'], 6)]},
                         'properties': props})
    for e in previas:
        lon, lat = puntos[e['estacion']]['geometry']['coordinates']
        features.append({'type': 'Feature',
                         'geometry': {'type': 'Point', 'coordinates': [round(lon, 6), round(lat, 6)]},
                         'properties': {
                             'grupo': 'exclusion_previa',
                             'rio': e['rio'], 'rio_display': rio_display(e['rio']),
                             'etiqueta': f"{rio_display(e['rio'])[4:]} {e['subtramo']}",
                             'subtramo': e['subtramo'],
                             'posicion_tabla_5_2': e['posicion_tabla_5_2'], 'categoria': e['categoria'],
                             'delimitacion': e['delimitacion'], 'area_cana_ha': round(e['area_cana_ha'], 2),
                             'estacion': e['estacion'], 'estado': 'Exclusión previa',
                             'motivo_exclusion': e['motivo_exclusion'],
                         }})
    fc = {
        'type': 'FeatureCollection',
        'metadata': {
            'titulo': 'Monitoreo del corredor: traducción de carga difusa a concentración (Actividad 5)',
            'generado': datetime.date.today().isoformat(),
            'generador': 'scripts/build_monitoreo_corredor.py',
            'factor_emision_n_kg_ha_ano': round(fe_n, 10),
            'factor_emision_p_kg_ha_ano': round(fe_p, 10),
            'retencion': RETENCION,
            'condicion_hidrologica': CONDICION_DEFECTO,
            'umbral_dominancia_puntual': UMBRAL_DOMINANCIA,
            'verificacion': 'docs/monitoreo_corredor.xlsx',
            'inventario_vertimientos': 'docs/Tasas retributivas.xlsx (CVC; DBO5/SST, sin N ni P)',
        },
        'features': features,
    }
    with open(GEOJSON_OUT, 'w', encoding='utf-8') as fh:
        json.dump(fc, fh, ensure_ascii=False, indent=1)
        fh.write('\n')


def escribir_config_js(fe_n, fe_p):
    """src/config/monitoreo.js: única copia de las constantes en el geovisor.
    Se genera aquí para que nunca diverja de la hoja Parámetros del Excel."""
    contenido = f"""/* monitoreo.js — Constantes del panel "Monitoreo de calidad de agua".
 *
 * GENERADO por scripts/build_monitoreo_corredor.py; no editar a mano. Son las
 * mismas cifras de la hoja Parámetros de docs/monitoreo_corredor.xlsx, de modo
 * que el geovisor y el Excel no pueden quedar con valores distintos.
 *
 *   FE_N = {DOSIS_N:g} kg N/ha/año × {FRAC_N:.4f} (Informe Fase 1, Tablas 4.6 y 4.7)
 *   FE_P = {DOSIS_P2O5:g} kg P2O5/ha/año × {CONV_P2O5_P} × {FRAC_P:.4f} (Tablas 4.6 y 4.8)
 */

export const FE_N = {round(fe_n, 10)!r};  // kg N/ha/año
export const FE_P = {round(fe_p, 10)!r};  // kg P/ha/año
export const DIAS_ANO = {DIAS_ANO};
export const CONV_Q = {CONV_Q!r};  // C (mg/L) = carga (kg/d) / (Q (m³/s) × 86,4)

/* Rango de remoción validado por el proyecto: escenario 2, corredor establecido. */
export const REMOCION_MAX_VALIDADA = {round(1 - RETENCION['esc2'], 10)!r};

/* Escenarios del Excel (remoción = 1 − factor de retención). */
export const ESCENARIOS = [
  {{ remocion: {round(1 - RETENCION['base'], 10)!r}, etiqueta: 'Base' }},
  {{ remocion: {round(1 - RETENCION['esc1'], 10)!r}, etiqueta: 'Esc. 1' }},
  {{ remocion: {round(1 - RETENCION['esc2'], 10)!r}, etiqueta: 'Esc. 2' }},
];

/* Condición hidrológica → propiedad del GeoJSON con el caudal de referencia. */
export const CONDICIONES = [
  {{ id: 'Promedio',   prop: 'q_promedio_m3s' }},
  {{ id: 'Verano',     prop: 'q_verano_m3s' }},
  {{ id: 'Transición', prop: 'q_transicion_m3s' }},
  {{ id: 'Invierno',   prop: 'q_invierno_m3s' }},
];
"""
    os.makedirs(os.path.dirname(CONFIG_JS_OUT), exist_ok=True)
    with open(CONFIG_JS_OUT, 'w', encoding='utf-8', newline='\n') as fh:
        fh.write(contenido)


# ── Principal ──────────────────────────────────────────────────────────
def main():
    for p in (POSICIONES_CSV, TABLA_4_10_CSV, TASAS_XLSX, EQUIV_TASAS_CSV):
        if not os.path.isfile(p):
            raise SystemExit(f'Falta el insumo:\n  {p}\n(ver cabecera del script)')

    prior, tramos, posiciones, tabla_4_10, puntos, hidro = cargar_insumos()
    ejercicio, previas, compartidas = seleccionar(prior, tramos, posiciones, puntos, hidro)

    nombres_hidro = sorted({e['estacion_hidro'] for e in ejercicio})
    series, caudales = {}, {}
    for n in nombres_hidro:
        s = serie_caudal(n)
        series[n] = s
        q, cnt = caudales_por_condicion(s, hidro[n]['umbral_verano_m3s'], hidro[n]['umbral_invierno_m3s'])
        caudales[n] = {'q': q, 'n': cnt}

    fe_n, fe_p = calcular(ejercicio, tabla_4_10, caudales)
    vertimientos = cargar_vertimientos(posiciones)
    listado_vert = asignar_vertimientos(ejercicio, vertimientos, posiciones)
    escribir_excel(ejercicio, previas, compartidas, tabla_4_10, series, hidro, fe_n, fe_p, listado_vert)
    estado_recalculo = recalcular_con_excel(XLSX_OUT)
    escribir_geojson(ejercicio, previas, puntos, fe_n, fe_p)
    escribir_config_js(fe_n, fe_p)

    # ── Resumen ────────────────────────────────────────────────────────
    print(f'Factor de emisión: N {fe_n:.4f} kg/ha/año · P {fe_p:.4f} kg/ha/año')
    print(f'Universo Alta/Muy alta: {len(ejercicio) + len(previas)} subtramos · '
          f'exclusión previa: {len(previas)} · en el ejercicio: {len(ejercicio)}')
    if compartidas:
        print('Estaciones de calidad compartidas:', compartidas)
    print()
    print(f"{'Subtramo':<18}{'Estación de calidad':<48}{'Hidrom.':<14}{'Q':>7}"
          f"{'CN base':>9}{'N hist':>8}{'cociente':>9}{'CP base':>9}{'P hist':>8}{'% punt.':>8}  Estado")
    for e in ejercicio:
        dom = '' if e['frac_puntual_n'] is None else f"{e['frac_puntual_n'] * 100:.1f}"
        print(f"{e['rio'] + ' ' + str(e['subtramo']):<18}{e['estacion'][:46]:<48}{e['estacion_hidro']:<14}"
              f"{e['q_ref']:>7.3f}{e['conc_n_base']:>9.4f}{e['nt']:>8.3f}{e['cociente_n']:>9.3f}"
              f"{e['conc_p_base']:>9.4f}{e['pt']:>8.3f}{dom:>8}  {e['estado']}")
    for e in previas:
        print(f"{e['rio'] + ' ' + str(e['subtramo']):<18}{e['motivo_exclusion']}")

    print('\nCarga puntual del inventario de tasas retributivas (subtramos pendientes):')
    print(f"{'Subtramo':<20}{'n':>3}{'Q (L/s)':>10}{'DBO5 kg/d':>12}{'SST kg/d':>12}{'N difusa kg/d':>15}  Criterio")
    for e in ejercicio:
        if e.get('tramo_porh'):
            continue
        print(f"{e['etiqueta']:<20}{e['n_vertimientos']:>3}{e['q_vertido_l_s']:>10.2f}"
              f"{e['carga_dbo5_kg_d']:>12.2f}{e['carga_sst_kg_d']:>12.2f}{e['carga_n_base']:>15.2f}  "
              f"{e['criterio_dominancia']}")
        for nota in e['notas_vertimientos']:
            print(f"{'':<20}· {nota}")
    print(f'\n→ {os.path.relpath(XLSX_OUT, PROJECT_DIR)}  ({estado_recalculo})')
    print(f'→ {os.path.relpath(GEOJSON_OUT, PROJECT_DIR)}')
    print(f'→ {os.path.relpath(CONFIG_JS_OUT, PROJECT_DIR)}')


if __name__ == '__main__':
    main()
